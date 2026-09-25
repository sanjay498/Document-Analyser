"""
Isolated Unit Tests for Dynamic Table Cells and Table Row Duplication
(Phase 2 Isolated Validation)
"""

from io import BytesIO
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from backend.app.core.doc_processor import (
    detect_yellow_highlights,
    duplicate_and_populate_table_rows,
    apply_field_values_to_template,
    is_run_yellow_highlighted
)


def create_template_with_dynamic_table() -> Document:
    """
    Creates an in-memory document with a dynamic items table.
    """
    doc = Document()
    doc.add_heading("Statement of Work & Milestones", level=1)

    table = doc.add_table(rows=2, cols=4)
    table.style = 'Table Grid'

    # Row 0: Header Row
    headers = ["Milestone #", "Deliverable Description", "Completion Date", "Milestone Fee"]
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"

    # Row 1: Dynamic Template Row (with yellow highlighted cells)
    row1 = table.rows[1]

    # Cell 0: Milestone #
    c0 = row1.cells[0].paragraphs[0].add_run("M-1")
    c0.font.highlight_color = WD_COLOR_INDEX.YELLOW
    c0.bold = True

    # Cell 1: Deliverable
    c1 = row1.cells[1].paragraphs[0].add_run("Cloud Architecture & Kubernetes Design")
    c1.font.highlight_color = WD_COLOR_INDEX.YELLOW
    c1.italic = True

    # Cell 2: Date
    c2 = row1.cells[2].paragraphs[0].add_run("2026-11-30")
    c2.font.highlight_color = WD_COLOR_INDEX.YELLOW

    # Cell 3: Fee
    c3 = row1.cells[3].paragraphs[0].add_run("$50,000 USD")
    c3.font.highlight_color = WD_COLOR_INDEX.YELLOW
    c3.bold = True

    return doc


def test_isolated_dynamic_row_duplication():
    """
    Tests duplicate_and_populate_table_rows in complete isolation.
    Duplicates 1 template row into 4 populated data rows.
    """
    doc = create_template_with_dynamic_table()
    table = doc.tables[0]

    records = [
        {
            "Milestone #": "M-1",
            "Deliverable Description": "Phase 1: Zero-Trust Network & IAM Architecture",
            "Completion Date": "October 30, 2026",
            "Milestone Fee": "$65,000 USD"
        },
        {
            "Milestone #": "M-2",
            "Deliverable Description": "Phase 2: Automated CI/CD & Terraform Pipeline",
            "Completion Date": "December 15, 2026",
            "Milestone Fee": "$85,000 USD"
        },
        {
            "Milestone #": "M-3",
            "Deliverable Description": "Phase 3: SOC2 Compliance & Penetration Testing",
            "Completion Date": "February 28, 2027",
            "Milestone Fee": "$45,000 USD"
        },
        {
            "Milestone #": "M-4",
            "Deliverable Description": "Phase 4: Production Cutover & 24/7 SRE Runbook",
            "Completion Date": "April 15, 2027",
            "Milestone Fee": "$55,000 USD"
        }
    ]

    # Run isolated row duplication
    duplicate_and_populate_table_rows(
        table=table,
        template_row_index=1,
        records=records,
        clear_highlight=True
    )

    # Verifications:
    # 1. Total rows = 1 header + 4 records = 5 rows
    assert len(table.rows) == 5, f"Expected 5 table rows, got {len(table.rows)}"

    # 2. Check content of each row
    for i, rec in enumerate(records):
        row = table.rows[i + 1]
        assert row.cells[0].text == rec["Milestone #"]
        assert row.cells[1].text == rec["Deliverable Description"]
        assert row.cells[2].text == rec["Completion Date"]
        assert row.cells[3].text == rec["Milestone Fee"]

        # Ensure yellow highlight is cleared on all cells
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    assert not is_run_yellow_highlighted(r), f"Highlight remained on row {i+1} cell"


def test_table_cell_context_detection():
    """
    Tests that detect_yellow_highlights captures column header and row context for table cells.
    """
    doc = create_template_with_dynamic_table()
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    parsed_doc, fields, table_groups = detect_yellow_highlights(bio.getvalue())

    # We should have 4 table cell fields
    assert len(fields) == 4
    assert len(table_groups) == 1

    group = table_groups[0]
    assert group.table_index == 0
    assert group.template_row_index == 1
    assert len(group.columns) == 4
    assert group.columns[0].header == "Milestone #"
    assert group.columns[1].header == "Deliverable Description"
    assert group.columns[2].header == "Completion Date"
    assert group.columns[3].header == "Milestone Fee"

    # Check cell fields have column headers and row context
    f_fee = next(f for f in fields if f.location.col_index == 3)
    assert f_fee.is_table_cell is True
    assert f_fee.column_header == "Milestone Fee"
    assert "Milestone #" in f_fee.row_context
    assert "[Column: 'Milestone Fee']" in f_fee.context_with_marker


def test_full_table_group_replacement_via_apply_field_values():
    """
    Tests applying dynamic table groups via apply_field_values_to_template.
    """
    doc = create_template_with_dynamic_table()
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    template_bytes = bio.getvalue()

    _, fields, table_groups = detect_yellow_highlights(template_bytes)

    table_records = {
        table_groups[0].group_id: [
            {"0": "Item A", "1": "Design Deliverable", "2": "Nov 1", "3": "$20,000"},
            {"0": "Item B", "1": "Build Deliverable", "2": "Dec 1", "3": "$30,000"},
            {"0": "Item C", "1": "Deploy Deliverable", "2": "Jan 1", "3": "$40,000"}
        ]
    }

    output_bio = apply_field_values_to_template(
        template_source=template_bytes,
        fields=fields,
        field_values={},
        table_group_records=table_records,
        clear_highlight=True
    )

    out_doc = Document(output_bio)
    t = out_doc.tables[0]
    assert len(t.rows) == 4  # 1 header + 3 items
    assert t.rows[1].cells[0].text == "Item A"
    assert t.rows[2].cells[0].text == "Item B"
    assert t.rows[3].cells[0].text == "Item C"
