"""
Unit tests verifying that all table data across templates (specifically Legal Opinion tables)
are accurately replaced with real document facts, with 0 dummy records (Doc No. 2480/2012)
and 0 OCR header leaks (--- [Document: ...]).
"""

import os
from io import BytesIO
import pytest
import pytest_asyncio
from docx import Document

from backend.app.core.samples import generate_legal_opinion_title_report_template
from backend.app.core.doc_processor import detect_yellow_highlights, apply_field_values_to_template
from backend.app.core.ai_extractor import mock_heuristic_extractor
from backend.app.core.source_extractor import extract_text_from_pdf, ExtractedSourceDocument
from backend.app.db.database import init_db


@pytest_asyncio.fixture(autouse=True)
async def prepare_database():
    await init_db()


def test_table_data_replacement_ganapathy():
    # 1. Load Ganapathy Sale Deed PDF
    pdf_path = "/Users/apple/.gemini/antigravity/brain/8be68963-6439-419d-b754-e35983734649/.user_uploaded/media_1789962935652.pdf"
    assert os.path.exists(pdf_path), "Ganapathy media_1789962935652.pdf must exist"

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    doc_obj = extract_text_from_pdf(pdf_bytes, "media_1789962935652.pdf")
    source_docs = [doc_obj]

    # 2. Load Legal Opinion template
    template_bio = generate_legal_opinion_title_report_template()
    template_bytes = template_bio.getvalue()

    _, fields, table_groups = detect_yellow_highlights(template_bytes)
    assert len(fields) > 10
    assert len(table_groups) >= 2

    # 3. Extract field values and table groups
    extraction_output = mock_heuristic_extractor(
        fields=fields,
        source_docs=source_docs,
        table_groups=table_groups
    )

    field_values = {f.field_id: f.value for f in extraction_output.fields if f.value is not None}
    table_group_records = {tg.group_id: tg.records for tg in extraction_output.table_groups}

    # 4. Generate populated .docx
    filled_bio = apply_field_values_to_template(
        template_source=template_bytes,
        fields=fields,
        field_values=field_values,
        table_group_records=table_group_records,
        clear_highlight=True
    )
    filled_doc = Document(filled_bio)

    # 5. Assertions on Tables:
    # Table 0: Scrutinized Documents
    t0 = filled_doc.tables[0]
    t0_text = " ".join(cell.text for row in t0.rows for cell in row.cells)
    assert "Doc No. 2480/2012" not in t0_text, "Dummy deed record 2480/2012 found in Table 0"
    assert "Pollachi SRO" not in t0_text, "Dummy Pollachi SRO found in Table 0"
    assert "1931/2026" in t0_text, "Sale deed 1931/2026 not found in Table 0"
    assert "2874/2018" in t0_text, "Prior deed 2874/2018 not found in Table 0"
    assert "Komangalam" in t0_text, "Komangalam SRO not found in Table 0"
    assert "--- [Document:" not in t0_text, "OCR header leak found in Table 0"

    # Table 1: Property Description
    t1 = filled_doc.tables[1]
    t1_text = " ".join(cell.text for row in t1.rows for cell in row.cells)
    assert "LAKSHMI" in t1_text, "Lakshmi not found in Table 1"
    assert "84/A2" in t1_text, "S.F.No. 84/A2 not found in Table 1"
    assert "1.28 Acres" in t1_text or "0.52.0" in t1_text, "1.28 Acres / 0.52.0 not found in Table 1"
    assert "Pannaikinaru" in t1_text, "Pannaikinaru Village not found in Table 1"
    assert "2.57 ACRES" not in t1_text, "Stale Muthulakshmi Item 2 found in Table 1"
    assert "--- [Document:" not in t1_text, "OCR header leak found in Table 1"

    # Table 2: Possession Certificate Table
    t2 = filled_doc.tables[2]
    t2_text = " ".join(cell.text for row in t2.rows for cell in row.cells)
    assert "84/A2" in t2_text, "84/A2 not found in Table 2"
    assert "0.52.0" in t2_text, "0.52.0 not found in Table 2"
    assert "--- [Document:" not in t2_text, "OCR header leak found in Table 2"
    assert "(OCR)" not in t2_text, "OCR marker leak found in Table 2"
    assert len(t2.rows) <= 5, f"Table 2 rows bloated: {len(t2.rows)}"

    # Table 3: Questions Table (37 questions)
    t3 = filled_doc.tables[3]
    t3_text = " ".join(cell.text for row in t3.rows for cell in row.cells)
    assert "--- [Document:" not in t3_text, "OCR header leak found in Table 3"
    assert "(OCR)" not in t3_text, "OCR marker leak found in Table 3"
    assert len(t3.rows) <= 45, f"Table 3 rows bloated: {len(t3.rows)}"

    # Table 4: Section 5 Documents to be obtained
    t4 = filled_doc.tables[4]
    t4_text = " ".join(cell.text for row in t4.rows for cell in row.cells)
    assert "Doc No. 2480/2012" not in t4_text, "Dummy deed record 2480/2012 found in Table 4"
    assert "1931/2026" in t4_text, "Sale deed 1931/2026 not found in Table 4"
    assert "Komangalam" in t4_text, "Komangalam SRO not found in Table 4"

    # Table 6: Annexure I Checklist Table
    t6 = filled_doc.tables[6]
    t6_text = " ".join(cell.text for row in t6.rows for cell in row.cells)
    assert "Udumalaipettai Branch" in t6_text, "Branch not replaced in Table 6"
    assert "LAKSHMI" in t6_text, "Lakshmi not found in Table 6"
    assert "Pannaikinaru Village, Udumalaipettai Taluk" in t6_text, "Location not replaced in Table 6"
    assert "84/A2" in t6_text, "Survey No 84/A2 not found in Table 6"
    assert "0.52.0 Hectare" in t6_text or "1.28 Acres" in t6_text, "Extent not replaced in Table 6"
    assert "04.06.2026" in t6_text, "Execution date 04.06.2026 not found in Table 6"
    assert "--- [Document:" not in t6_text, "OCR header leak found in Table 6"


def test_table_data_replacement_balashanmugam():
    doc_text = (
        "REGISTERED PARTITION DEED\n"
        "Document No. 1773/1998, SRO Anaimalai, dated 08.10.1998.\n"
        "Settlement Deed Doc No. 5035/2012, SRO Anaimalai.\n"
        "Borrower / Title Holder: Balashanmugam, S/o Kalimuthu Chettiyar\n"
        "Power Agent: Senthilraja, S/o Balashanmugam\n"
        "Property situated at Thensangampalayam Village, Pollachi Taluk, Coimbatore District.\n"
        "Survey Numbers: S.F.No. 74/B, 75, and 76/2.\n"
        "Total Extent: 6.11 Acres (S.F.74/B: 3.73 Acres, S.F.75: 0.64 Acres, S.F.76/2: 1.74 Acres).\n"
        "Boundaries: North of East-West Main Road, South of Lands in S.F.No. 83, 76/2, and 75, "
        "East of Lands belonging to Venugopal and S.F.No. 93, West of Lands belonging to Arumugam in S.F.No. 75.\n"
        "Encumbrance Certificate for 30 years discloses Nil encumbrance.\n"
    )
    source_docs = [ExtractedSourceDocument(
        filename="balashanmugam_deed.txt",
        file_type="txt",
        full_text=doc_text,
        page_texts=[doc_text],
        char_count=len(doc_text),
        page_or_section_count=1,
        is_ocr=False
    )]

    template_bio = generate_legal_opinion_title_report_template()
    template_bytes = template_bio.getvalue()

    _, fields, table_groups = detect_yellow_highlights(template_bytes)
    extraction_output = mock_heuristic_extractor(
        fields=fields,
        source_docs=source_docs,
        table_groups=table_groups
    )

    field_values = {f.field_id: f.value for f in extraction_output.fields if f.value is not None}
    table_group_records = {tg.group_id: tg.records for tg in extraction_output.table_groups}

    filled_bio = apply_field_values_to_template(
        template_source=template_bytes,
        fields=fields,
        field_values=field_values,
        table_group_records=table_group_records,
        clear_highlight=True
    )
    filled_doc = Document(filled_bio)

    # Table 0: Balashanmugam deeds
    t0_text = " ".join(cell.text for row in filled_doc.tables[0].rows for cell in row.cells)
    assert "Doc No. 2480/2012" not in t0_text
    assert "1773/1998" in t0_text
    assert "5035/2012" in t0_text
    assert "Anaimalai" in t0_text

    # Table 1: Balashanmugam property
    t1_text = " ".join(cell.text for row in filled_doc.tables[1].rows for cell in row.cells)
    assert "Balashanmugam" in t1_text
    assert "6.11 Acres" in t1_text
    assert "Thensangampalayam" in t1_text
    assert "74/B" in t1_text

    # Table 6: Balashanmugam checklist
    t6_text = " ".join(cell.text for row in filled_doc.tables[6].rows for cell in row.cells)
    assert "Balashanmugam" in t6_text
    assert "Thensangampalayam" in t6_text
    assert "6.11 Acres" in t6_text
