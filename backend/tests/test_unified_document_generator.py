import pytest
import io
import json
from io import BytesIO
from docx import Document
from httpx import AsyncClient, ASGITransport

from backend.app.main import app
from backend.app.core.qa_engine import extract_template_questions
from backend.app.core.deed_models import format_certificate_of_title, generate_multi_paragraph_trace


def create_unified_template_docx() -> bytes:
    doc = Document()
    doc.add_heading("BANK LEGAL OPINION & SCRUTINY", level=1)
    p = doc.add_paragraph("Borrower Name: ")
    # Add a yellow highlighted field
    r = p.add_run("M.Anandan, S/o.Mayilsamy Kavundar")
    from docx.enum.text import WD_COLOR_INDEX
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_heading("2. TITLE SCRUTINY QUESTIONS", level=2)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'
    tbl.rows[0].cells[0].text = "Scrutiny Query"
    tbl.rows[0].cells[1].text = "Findings / Compliance"

    r1 = tbl.add_row().cells
    r1[0].text = "Whether the chain of title for 30 years is complete without missing links?"
    r1[1].text = ""

    r2 = tbl.add_row().cells
    r2[0].text = "Whether there are any prior subsisting encumbrances or mortgages over the property?"
    r2[1].text = ""

    doc.add_heading("3. CERTIFICATE OF TITLE", level=2)
    p_cert = doc.add_paragraph(
        "(Original fee receipts enclosed). I certify that M.Anandan, S/o.Mayilsamy Kavundar has an absolute, clear and "
        "marketable title over the property situated at Mannur Village and offered for mortgage."
    )
    p_cert.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


@pytest.mark.asyncio
async def test_unified_document_generator_end_to_end():
    from backend.app.db.database import init_db
    await init_db()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Create a session
        create_resp = await client.post("/api/sessions")
        assert create_resp.status_code == 200
        session_id = create_resp.json()["session_id"]

        # 2. Upload template to single intake desk
        template_bytes = create_unified_template_docx()
        upload_tpl_resp = await client.post(
            f"/api/sessions/{session_id}/template",
            files={"file": ("Unified_Scrutiny_Template.docx", template_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        assert upload_tpl_resp.status_code == 200
        tpl_data = upload_tpl_resp.json()
        assert tpl_data["fields_count"] >= 1
        assert tpl_data["questions_count"] >= 2

        # 3. Upload source title deed
        source_doc_text = """
        SALE DEED (Doc No: 1931/2026)
        Sub-Registrar Office: Komangalam
        Vendor: C. Ganapathy, S/o Chinnan
        Purchaser / Title Holder: V. LAKSHMI, W/o Vellingiri
        Property: S.F.No.84/A2, Pannaikinaru Village, Extent: 1.28 Acres (0.52.0 Hectare).
        Prior Title: S.F.No.84/A originally purchased under Doc No: 2874/2018.
        Encumbrance Certificate search from 1996 to 2026 shows Nil encumbrances.
        Patta No. 2335 issued in the name of the owner with continuous possession.
        """
        source_bio = BytesIO(source_doc_text.encode("utf-8"))
        upload_src_resp = await client.post(
            f"/api/sessions/{session_id}/sources",
            files={"files": ("Doc_1931_2026_Sale_Deed.txt", source_bio, "text/plain")}
        )
        assert upload_src_resp.status_code == 200

        # 4. Run unified extraction (Both highlighted fields and grounded questions simultaneously!)
        extract_resp = await client.post(f"/api/sessions/{session_id}/extract")
        assert extract_resp.status_code == 200
        ext_data = extract_resp.json()
        assert ext_data["total_fields"] >= 1
        assert ext_data["questions_count"] >= 2
        assert len(ext_data["qa_answers"]) >= 2

        # Verify Q&A answers have evidence citations
        qa_ans = ext_data["qa_answers"]
        chain_q = next((a for a in qa_ans if "chain of title" in a["question_text"].lower()), None)
        assert chain_q is not None
        assert chain_q["status"] in ("grounded", "answered", "partial")

        # 5. Document Rename API
        custom_name = "Legal_Opinion_Lakshmi_Vellingiri.docx"
        rename_resp = await client.post(
            f"/api/sessions/{session_id}/rename",
            json={"filename": custom_name}
        )
        assert rename_resp.status_code == 200
        assert rename_resp.json()["filename"] == custom_name

        # 6. Export Unified Document with field values and Q&A answers
        field_vals = {}
        for f in ext_data["results"]:
            field_vals[f["field_id"]] = f["value"] or "V. LAKSHMI, W/o Vellingiri"

        export_resp = await client.post(
            f"/api/sessions/{session_id}/export",
            json={
                "field_values": field_vals,
                "qa_answers": ext_data["qa_answers"],
                "doc_custom_name": custom_name
            }
        )
        assert export_resp.status_code == 200
        exp_data = export_resp.json()
        assert exp_data["status"] == "success"

        # 7. Download document and verify custom filename header
        down_resp = await client.get(f"/api/sessions/{session_id}/download")
        assert down_resp.status_code == 200
        assert "Legal_Opinion_Lakshmi_Vellingiri.docx" in down_resp.headers.get("Content-Disposition", "")

        # 8. Inspect generated docx bytes: ensure answers are injected into tables
        gen_doc = Document(BytesIO(down_resp.content))
        assert len(gen_doc.tables) >= 1
        table_cells_text = " ".join([c.text for row in gen_doc.tables[0].rows for c in row.cells])
        # Table should have answers populated instead of empty cells
        assert len(table_cells_text.strip()) > len("Scrutiny Query Findings / Compliance")

        # 9. Save as Next Template API
        save_tpl_resp = await client.post(
            f"/api/sessions/{session_id}/save-as-template",
            json={"name": "Lakshmi Bank Template", "bank_name": "SBI"}
        )
        assert save_tpl_resp.status_code == 200
        assert save_tpl_resp.json()["status"] == "success"
        assert save_tpl_resp.json()["name"] == "Lakshmi Bank Template"


def test_trace_no_price_and_certificate_clean():
    # 1. Ensure trace never includes price of land
    trace_paras = generate_multi_paragraph_trace(
        "sale_deed",
        context={
            "purchaser": "Balashanmugam, S/o Kalimuthu Chettiyar",
            "sf_nos": "74/B, 75",
            "extent": "6.11 Acres",
            "village": "Thensangampalayam Village",
            "sro": "Anaimalai",
            "date": "08.10.1998",
            "doc_no": "1773",
            "year": "1998"
        },
        paragraph_count=4
    )
    for p in trace_paras:
        assert "Rs." not in p
        assert "INR" not in p
        assert "for a sum of" not in p
        assert "price" not in p.lower()

    # 2. Ensure certificate deduplication and clean dummy name replacement
    raw_cert = (
        "The title holder Balashanmugam, S/o Kalimuthu Chettiyar holds absolute, clear, and marketable title over the\n"
        "properties and is legally competent to create mortgage security.\n"
        "(Original fee receipts enclosed). I certify that M.Anandan, S/o.Mayilsamy Kavundar has an absolute, clear and\n"
        "marketable title over the property.\n"
        "The title holder Balashanmugam, S/o Kalimuthu Chettiyar holds absolute, clear, and marketable title over the\n"
        "properties and is legally competent to create mortgage security."
    )
    cleaned = format_certificate_of_title(raw_cert, {"borrower": "Balashanmugam, S/o Kalimuthu Chettiyar"})
    assert "Balashanmugam, S/o Kalimuthu Chettiyar" in cleaned
    assert "M.Anandan" not in cleaned
    assert "Mayilsamy Kavundar" not in cleaned
    # Ensure redundant trace filler sentence is completely eliminated
    assert cleaned.count("The title holder Balashanmugam") == 0


@pytest.mark.asyncio
async def test_checklist_table_cells_never_missing():
    """Ensure that all table checklist items (Borrower, Extent, Survey No, Location, Taxes, Boundaries) are 100% recognized with 0 missing."""
    from backend.app.db.database import init_db
    from docx.enum.text import WD_COLOR_INDEX
    await init_db()

    # Create a template with checklist table
    doc = Document()
    doc.add_heading("CHECKLIST LEGAL SCRUTINY", level=1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    tbl.rows[0].cells[0].text = "Sr.No."
    tbl.rows[0].cells[1].text = "Particulars"
    tbl.rows[0].cells[2].text = "Compliance"

    items = [
        ("1.", "Name of the Branch", "Pollachi Branch"),
        ("2.", "Name of the Borrower", "Dummy Borrower"),
        ("3.", "Extent of area (in acres/sq.ft.)", "4.57 Acres"),
        ("4.", "Survey no/Gut no/CST no.", "S.F.No.245/1B"),
        ("5.", "Boundaries", "Details mentioned in separate sheet"),
        ("6.", "Location", "Mannur Village"),
        ("7.", "Taxes paid up to date", "Old tax receipt"),
        ("8.", "Type of land", "Agricultural"),
    ]
    for sr, part, comp in items:
        row = tbl.add_row().cells
        row[0].text = sr
        row[1].text = part
        p = row[2].paragraphs[0]
        r = p.add_run(comp)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    bio = BytesIO()
    doc.save(bio)
    template_bytes = bio.getvalue()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Create session
        create_resp = await client.post("/api/sessions")
        session_id = create_resp.json()["session_id"]

        # 2. Upload template
        await client.post(
            f"/api/sessions/{session_id}/template",
            files={"file": ("Checklist_Template.docx", template_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )

        # 3. Upload deed with Anandan / 711 / Kottur / 2223 Sq.ft. / Property tax receipt 5902
        source_text = """
        REGISTERED SALE DEED (Doc No. 750/1998 at SRO Anaimalai)
        Vendor: Rathinasamy Gounder
        Purchaser / Borrower: M.Anandan, S/o Mayilsamy Kavundar
        Property: S.F.No. 711 (New S.F.No. 711/2B2), Kottur Village, Anaimalai Taluk.
        Extent: 2223 Sq.ft. Residential House site.
        Boundaries: North by Rathinasamy Property, South by Senniyappa Gounder House, East by 30 Feet Road, West by North-South Road.
        Property Tax Receipt No. 5902 for year 2025-2026 paid up to date in name of M.Anandan.
        """
        await client.post(
            f"/api/sessions/{session_id}/sources",
            files={"files": ("Deed_Anandan.txt", BytesIO(source_text.encode("utf-8")), "text/plain")}
        )

        # 4. Extract
        extract_resp = await client.post(f"/api/sessions/{session_id}/extract")
        assert extract_resp.status_code == 200
        ext_data = extract_resp.json()

        # Check results
        assert ext_data["not_found_count"] == 0
        assert ext_data["extracted_count"] == len(items)
        assert ext_data["total_fields"] == len(items)

        res_by_part = {}
        for r in ext_data["results"]:
            assert r["status"] == "extracted"
            assert r["value"] is not None
            assert len(r["value"].strip()) > 0
