"""
Test Suite: In-Browser Yellow Highlighter Studio & Document Compiler API.
"""

import pytest
import pytest_asyncio
from starlette.testclient import TestClient

from backend.app.main import app
from backend.app.db.database import init_db
from backend.app.api.editor_routes import (
    build_docx_from_editor_payload,
    EditorDocumentPayload,
    ParagraphPayload,
    TextRunPayload,
    DynamicTablePayload,
    TableColumnPayload
)
from backend.app.core.doc_processor import detect_yellow_highlights


@pytest_asyncio.fixture(autouse=True)
async def prepare_database():
    await init_db()


def test_build_docx_from_editor_payload_with_highlights():
    payload = EditorDocumentPayload(
        title="Custom_Legal_Opinion_Studio.docx",
        paragraphs=[
            ParagraphPayload(
                heading_level=1,
                alignment="center",
                runs=[TextRunPayload(text="TITLE SEARCH REPORT", is_highlighted=False, bold=True)]
            ),
            ParagraphPayload(
                alignment="left",
                runs=[
                    TextRunPayload(text="Borrower Name: ", is_highlighted=False, bold=True),
                    TextRunPayload(text="K.MUTHULAKSHMI, W/o G.Kumar", is_highlighted=True),
                    TextRunPayload(text=" | Loan Branch: ", is_highlighted=False),
                    TextRunPayload(text="Pollachi Branch", is_highlighted=True),
                ]
            ),
            ParagraphPayload(
                alignment="left",
                runs=[
                    TextRunPayload(text="The total land extent is ", is_highlighted=False),
                    TextRunPayload(text="4.57 Acres", is_highlighted=True),
                    TextRunPayload(text=" located in S.F.No. 245/1B and 245/3A2.", is_highlighted=False),
                ]
            )
        ],
        tables=[
            DynamicTablePayload(
                title="Document Scrutiny Table",
                columns=[
                    TableColumnPayload(header="S.No", sample_text="1", is_highlighted=False),
                    TableColumnPayload(header="Deed Type", sample_text="Sale Deed", is_highlighted=True),
                    TableColumnPayload(header="Doc No", sample_text="2480/2012", is_highlighted=True),
                    TableColumnPayload(header="SRO", sample_text="Pollachi SRO", is_highlighted=True),
                ]
            )
        ]
    )

    bio = build_docx_from_editor_payload(payload)
    raw_bytes = bio.getvalue()
    assert len(raw_bytes) > 1500

    # Verify that doc_processor detects the yellow highlights
    doc, fields, table_groups = detect_yellow_highlights(raw_bytes)
    assert len(fields) >= 3
    assert len(table_groups) == 1

    orig_texts = [f.original_text for f in fields]
    assert "K.MUTHULAKSHMI, W/o G.Kumar" in orig_texts
    assert "Pollachi Branch" in orig_texts
    assert "4.57 Acres" in orig_texts

    assert table_groups[0].table_index == 0
    assert len(table_groups[0].columns) == 4


def test_editor_api_endpoints():
    client = TestClient(app)

    payload = {
        "title": "In_Browser_Crafted_Contract.docx",
        "paragraphs": [
            {
                "heading_level": 1,
                "alignment": "center",
                "runs": [{"text": "CONSULTING CONTRACT", "is_highlighted": False, "bold": True}]
            },
            {
                "alignment": "left",
                "runs": [
                    {"text": "Contract Effective Date: ", "is_highlighted": False},
                    {"text": "November 01, 2026", "is_highlighted": True},
                    {"text": " between Acme and ", "is_highlighted": False},
                    {"text": "Globex Corp", "is_highlighted": True}
                ]
            }
        ],
        "tables": [
            {
                "title": "Payment Milestones",
                "columns": [
                    {"header": "Milestone", "sample_text": "M1", "is_highlighted": False},
                    {"header": "Amount", "sample_text": "$50,000", "is_highlighted": True}
                ]
            }
        ]
    }

    # 1. Test convert-to-docx download
    res_conv = client.post("/api/editor/convert-to-docx", json=payload)
    assert res_conv.status_code == 200
    assert len(res_conv.content) > 1000
    assert res_conv.headers["content-disposition"] == 'attachment; filename="In_Browser_Crafted_Contract.docx"'

    # 2. Test save-as-template
    res_save = client.post("/api/editor/save-as-template", json=payload)
    assert res_save.status_code == 200
    tpl_data = res_save.json()
    assert tpl_data["id"]
    assert tpl_data["name"] == "In_Browser_Crafted_Contract.docx"
    assert tpl_data["fields_count"] >= 2
    assert tpl_data["table_groups_count"] == 1

    # 3. Test use-in-session
    res_use = client.post("/api/editor/use-in-session", json=payload)
    assert res_use.status_code == 200
    use_data = res_use.json()
    assert use_data["session_id"]
    assert use_data["template_filename"] == "In_Browser_Crafted_Contract.docx"
    assert use_data["fields_count"] >= 2
    assert len(use_data["fields"]) >= 2
    assert len(use_data["table_groups"]) == 1

    # 4. Test import-template by ID
    template_id = tpl_data["id"]
    res_import_tpl = client.get(f"/api/editor/import-template/{template_id}")
    assert res_import_tpl.status_code == 200
    imported_payload = res_import_tpl.json()
    assert imported_payload["title"] == "In_Browser_Crafted_Contract.docx"
    assert len(imported_payload["paragraphs"]) >= 2
    assert len(imported_payload["tables"]) == 1

    # 5. Test update existing template
    updated_payload = imported_payload
    updated_payload["title"] = "Updated_Contract_Studio.docx"
    updated_payload["paragraphs"][1]["runs"].append({
        "text": " Extra highlighted term",
        "is_highlighted": True
    })
    res_update = client.put(f"/api/editor/templates/{template_id}", json=updated_payload)
    assert res_update.status_code == 200
    updated_data = res_update.json()
    assert updated_data["name"] == "Updated_Contract_Studio.docx"
    assert updated_data["fields_count"] >= 3

    # Clean up template
    client.delete(f"/api/templates/{template_id}")

    # 6. Test import file via upload
    from backend.app.core.samples import generate_legal_opinion_title_report_template
    legal_bio = generate_legal_opinion_title_report_template()
    res_upload_import = client.post(
        "/api/editor/import-file",
        files={"file": ("Legal_Report.docx", legal_bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res_upload_import.status_code == 200
    uploaded_payload = res_upload_import.json()
    assert uploaded_payload["title"] == "Legal_Report.docx"
    assert len(uploaded_payload["paragraphs"]) > 5
    # Verify yellow highlights preserved
    all_runs = [r for p in uploaded_payload["paragraphs"] for r in p["runs"]]
    assert any(r["is_highlighted"] for r in all_runs)
    assert len(uploaded_payload.get("elements", [])) > 0


def test_interleaved_tables_and_elements_preservation():
    import io
    import docx
    from backend.app.api.editor_routes import parse_docx_to_editor_payload

    doc = docx.Document()
    doc.add_paragraph("Section 1 Intro")
    tbl = doc.add_table(rows=2, cols=3)
    tbl.rows[0].cells[0].text = "S.No"
    tbl.rows[0].cells[1].text = "Deed"
    tbl.rows[0].cells[2].text = "Date"
    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "Sale Deed"
    tbl.rows[1].cells[2].text = "05.05.1987"
    doc.add_paragraph("Section 2 Detailed Trace")

    buf = io.BytesIO()
    doc.save(buf)

    payload = parse_docx_to_editor_payload(buf.getvalue(), "interleaved.docx")
    assert len(payload.elements) == 3
    assert payload.elements[0].type == "paragraph"
    assert payload.elements[1].type == "table"
    assert payload.elements[2].type == "paragraph"

    # Verify building docx maintains exact interleaved positions
    out_bio = build_docx_from_editor_payload(payload)
    rebuilt_payload = parse_docx_to_editor_payload(out_bio.getvalue(), "rebuilt.docx")
    assert len(rebuilt_payload.elements) == 3
def test_pdf_upload_table_preservation():
    """
    Verifies that uploading a PDF document preserves all structured tables and paragraphs
    without collapsing table rows into flat text strings.
    """
    import os
    import sqlite3
    import io
    from backend.app.api.editor_routes import parse_pdf_to_editor_payload

    db_path = os.getenv("DOCFILLER_DB_PATH", "test_docfiller.db")
    con = sqlite3.connect(db_path)
    row = con.execute("SELECT template_bytes FROM generation_sessions WHERE template_filename LIKE '%Muthulakshmi%' LIMIT 1").fetchone()
    if not row or not row[0] or not row[0].startswith(b"%PDF"):
        return

    pdf_bytes = row[0]
    payload = parse_pdf_to_editor_payload(pdf_bytes, "Muthulakshmi_Bank_Opinion.pdf")

    assert payload.title == "Muthulakshmi_Bank_Opinion.docx"
    assert len(payload.elements) > 10
    
    # Verify tables are preserved as structured table elements
    table_elements = [e for e in payload.elements if e.type == "table" and e.table]
    assert len(table_elements) >= 5, f"Expected at least 5 structured tables, found {len(table_elements)}"
    
    # Verify Table 1 has 6 columns
    first_tbl = table_elements[0].table
    assert len(first_tbl.columns) == 6
    assert "Sr. No." in first_tbl.columns[0].header
    assert len(first_tbl.rows) >= 4

    # Verify Paragraphs retain headers and alignments
    headings = [e.paragraph for e in payload.elements if e.type == "paragraph" and e.paragraph and e.paragraph.heading_level]
    assert len(headings) >= 3


def test_in_place_docx_xml_substitution_fidelity():
    """
    Verifies that in_place_docx_xml_substitute directly modifies word/document.xml
    without touching un-highlighted paragraphs, fonts, character spacing, or table layouts.
    """
    import io
    import docx
    import zipfile
    from backend.app.core.doc_processor import in_place_docx_xml_substitute

    doc = docx.Document()
    p1 = doc.add_paragraph("Static Unchanged Heading")
    p2 = doc.add_paragraph()
    r1 = p2.add_run("Fixed Prefix: ")
    r2 = p2.add_run("Original Highlighted Variable")
    r2.font.highlight_color = 7 # Yellow
    r3 = p2.add_run(" Fixed Suffix.")

    buf = io.BytesIO()
    doc.save(buf)
    tpl_bytes = buf.getvalue()

    replacements = {"Original Highlighted Variable": "New Substituted Legal Fact"}
    out_bytes = in_place_docx_xml_substitute(tpl_bytes, replacements, clear_highlight=True)

    with zipfile.ZipFile(io.BytesIO(out_bytes), 'r') as zout:
        doc_xml = zout.read("word/document.xml").decode("utf-8")
        assert "New Substituted Legal Fact" in doc_xml
        assert "Original Highlighted Variable" not in doc_xml
        assert "Static Unchanged Heading" in doc_xml
        assert "Fixed Prefix: " in doc_xml
        assert " Fixed Suffix." in doc_xml
        assert 'w:val="yellow"' not in doc_xml
